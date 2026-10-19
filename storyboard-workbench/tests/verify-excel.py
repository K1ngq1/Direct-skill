"""Independent readback: python verify-excel.py <project.json> <download.xlsx>."""
import json
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook


def normalized(value):
    return value.replace('\r\n', '\n').replace('\r', '\n')


def records(sheet):
    headers = [cell.value for cell in sheet[1]]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = {}
        for header, value in zip(headers, values):
            if '（续' in header:
                base = header.rsplit('（续', 1)[0]
                row[base] = (row[base] or '') + (value or '')
            else:
                row[header] = value
        rows.append(row)
    return rows


source = json.loads(Path(sys.argv[1]).read_text(encoding='utf-8-sig'))
filename = Path(sys.argv[2])
with zipfile.ZipFile(filename) as archive:
    assert 'xl/workbook.xml' in archive.namelist()
workbook = load_workbook(filename, data_only=False)
assert workbook.sheetnames == ['分镜总览', '总共的Prompt']
overview, prompts = [records(sheet) for sheet in workbook]
shots = source['shots']
assert shots and len(overview) == len(shots)
draft = source.get('workflowMode') == 'agent-default'
chosen = {}
for shot in shots:
    variant = next((v for v in shot['variants'] if v['id'] == shot.get('selectedVariantId')), None)
    if variant is None and draft:
        variant = next((v for v in shot['variants'] if v['id'] == shot.get('recommendation', {}).get('variantId')), None)
    assert variant is not None, ('missing actual choice or default', shot['id'])
    chosen[shot['id']] = variant
variant_column = '当前方案ID' if draft else '已选方案ID'
allowed_sections = [section for section in source.get('promptSections', []) if
                    (section.get('variantId') and section.get('shotIds') and
                     all(chosen[shot_id]['id'] == section['variantId'] for shot_id in section['shotIds'])) or
                    (not section.get('variantId') and section.get('category') != 'shotFrames')]
expected_count = len(shots) + sum(len(section.get('prompts', [])) for section in allowed_sections)
assert len(prompts) == expected_count, (len(prompts), expected_count)
for shot, row in zip(shots, overview):
    number = str(shot.get('displayNumber', shot.get('sourceNumber', shot['id'])))
    assert row['镜号'] == number and isinstance(row['镜号'], str)
    assert row[variant_column] == chosen[shot['id']]['id']
    assert row['建议秒数'] == shot.get('suggestedDuration')
    assert row['脚本画面工作稿'] == normalized(shot.get('workingScript', shot.get('original', '')))
    if draft:
        variant = chosen[shot['id']]
        selected = shot.get('selectedVariantId') == variant['id']
        approved = shot.get('approval') or {}
        status = '用户已确认当前版本' if selected and approved.get('variantId') == variant['id'] and approved.get('variantVersion') == variant['version'] else '用户已选择 · 待确认' if selected else 'Agent 默认方案 · 待审'
        assert row['方案来源'] == ('用户选择' if selected else 'Agent 默认方案')
        assert row['审核状态'] == status + ('；待同步' if variant.get('promptNeedsSync') else '')
        if shot.get('sourceNumber') is None:
            assert row['原镜号'] in (None, '')
for shot, row in zip(shots, prompts):
    assert row['类型'] == '分镜主静帧'
    assert row['关联镜号'] == str(shot.get('displayNumber', shot.get('sourceNumber', shot['id'])))
    assert row[variant_column] == chosen[shot['id']]['id']
    assert row['静态Prompt'] == normalized(chosen[shot['id']]['staticPrompt'])
expected_sections = [normalized(prompt['text']) for section in allowed_sections for prompt in section.get('prompts', [])]
assert [row['静态Prompt'] for row in prompts[len(shots):]] == expected_sections
for sheet in workbook:
    assert sheet.freeze_panes
    assert len(sheet.title) <= 31
    for row in sheet:
        for cell in row:
            if cell.value is not None:
                assert cell.data_type != 'f'
                assert cell.alignment.wrap_text
                if isinstance(cell.value, str):
                    assert len(cell.value.encode('utf-16-le')) // 2 <= 32767
print(json.dumps({'file': filename.name, 'sheets': workbook.sheetnames,
                  'mode': 'agent-default' if draft else 'manual', 'overview_rows': len(overview),
                  'draft_main_prompts' if draft else 'selected_main_prompts': len(shots),
                  'independent_and_draft_extra_prompts' if draft else 'independent_and_selected_extra_prompts': len(expected_sections),
                  'result': 'PASS'}, ensure_ascii=False))

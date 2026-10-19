import io
import json
import subprocess
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image

from test_import import module, PNG, ROOT


class MappingImportTest(unittest.TestCase):
    def setUp(self):
        artifacts = ROOT.parent / "artifacts"
        artifacts.mkdir(exist_ok=True)
        self.folder = Path(tempfile.mkdtemp(prefix="excel-mapping-verification-", dir=artifacts))
        self.source = self.folder / "映射输入.xlsx"

    def test_missing_numbers_merged_rows_and_images_keep_source(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "脚本"
        sheet.append(["创作说明"])
        sheet.append([])
        sheet.append([None, None, "镜头想法", "秒", "备注"])
        sheet.append([None, None, "人物躺在床上", "4.5秒", "抬着手机"])
        sheet.append([None, None, None, None, "手与手机同时可见"])
        sheet.append([None, None, "妈妈站在门口", 3, "拿着锅铲"])
        sheet.merge_cells("C4:C5")
        sheet.add_image(Image(io.BytesIO(PNG)), "B5")
        other = workbook.create_sheet("附表")
        other["A1"] = "另一个未转换的内容块"
        workbook.save(self.source)
        original = self.source.read_bytes()
        mapping = {"sheets": [{"sheet": "脚本", "headerRow": 3, "contentColumn": "c", "durationColumn": 4}]}
        result = module.import_excel(self.source, self.folder / "project", mapping)
        repeat = module.import_excel(self.source, self.folder / "repeat", mapping)
        self.assertEqual([shot["id"] for shot in result["shots"]], ["S001", "S002"])
        self.assertEqual(result["shots"], repeat["shots"])
        shot = result["shots"][0]
        self.assertEqual(shot["sourceNumber"], "")
        self.assertIsNone(shot["sourceRef"]["numberCell"])
        self.assertEqual(shot["sourceRef"]["contentRange"], "C4:C5")
        self.assertEqual(shot["originalDuration"], 4.5)
        self.assertEqual(shot["durationOrigin"], "source")
        self.assertIn({"name": "备注", "value": "手与手机同时可见", "cell": "E5"}, shot["originalFields"])
        self.assertEqual(result["references"][0]["shotId"], "S001")
        self.assertEqual(result["references"][0]["sha256"], module.digest(PNG))
        self.assertTrue(any("附表" in issue["text"] for issue in result["reviewIssues"]))
        self.assertEqual(len(result["source"]["sheets"]), 2)
        self.assertEqual(self.source.read_bytes(), original)
        self.assertEqual((self.folder / "project/sources/source.xlsx").read_bytes(), original)
        markdown = (self.folder / "project/sources/source.md").read_text(encoding="utf-8")
        self.assertIn("内部镜头 S001（原稿未提供镜号）", markdown)
        self.assertNotIn("原镜号 S001", markdown)

    def test_multiple_sheets_headerless_mapping_and_ambiguous_duration(self):
        workbook = openpyxl.Workbook()
        workbook.active.title = "甲"
        workbook.active.append(["01", "人物看向门口", "00:03"])
        workbook.create_sheet("乙").append([None, "空房间", "2s"])
        workbook.save(self.source)
        mapping = {"sheets": [
            {"sheet": "乙", "headerRow": 0, "contentColumn": 2, "numberColumn": 1, "durationColumn": 3},
            {"sheet": "甲", "headerRow": 0, "contentColumn": "B", "numberColumn": "A", "durationColumn": "C"},
        ]}
        result = module.import_excel(self.source, self.folder / "project", mapping)
        self.assertEqual([shot["sourceNumber"] for shot in result["shots"]], ["01", ""])
        self.assertEqual([shot["original"] for shot in result["shots"]], ["人物看向门口", "空房间"])
        self.assertIsNone(result["shots"][0]["originalDuration"])
        self.assertEqual(result["shots"][1]["originalDuration"], 2)
        self.assertTrue(any(issue["id"] == "duration-S001" for issue in result["reviewIssues"]))

    def test_formula_caches_are_used_and_missing_cache_is_not_prompt_text(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["分镜号", "画面内容", "时长", "备注"])
        sheet.append(["=10+2", '=CONCAT("房门","打开")', "=4+0.25", '=CONCAT("站在","门口")'])
        sheet.append(["=1+2", '=CONCAT("人物","抬眼")', "=5", "注释"])
        sheet.append([4, "=1/0", 2, None])
        workbook.save(self.source)
        ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        with zipfile.ZipFile(self.source) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        tree = ET.fromstring(files["xl/worksheets/sheet1.xml"])
        for coordinate, value, kind in [("A2", "12", "n"), ("B2", "房门打开", "str"), ("C2", "4.25", "n"), ("D2", "站在门口", "str"), ("B4", "#DIV/0!", "e")]:
            cell = tree.find(f"m:sheetData/m:row/m:c[@r='{coordinate}']", ns)
            cell.set("t", kind)
            cell.find("m:v", ns).text = value
        files["xl/worksheets/sheet1.xml"] = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
        with zipfile.ZipFile(self.source, "w") as archive:
            for name, data in files.items():
                archive.writestr(name, data)
        mapping = {"sheets": [{"sheet": "Sheet", "headerRow": 1, "contentColumn": "B", "numberColumn": "A", "durationColumn": "C"}]}
        result = module.import_excel(self.source, self.folder / "project", mapping)
        first, missing, error = result["shots"]
        self.assertEqual(first["sourceNumber"], "12")
        self.assertEqual(first["original"], "房门打开")
        self.assertEqual(first["originalDuration"], 4.25)
        self.assertIn({"name": "备注", "value": "站在门口", "cell": "D2"}, first["originalFields"])
        self.assertEqual(missing["original"], "")
        self.assertEqual(missing["sourceNumber"], "")
        self.assertEqual(error["original"], "")
        archived = result["source"]["sheets"][0]["cells"]["B2"]
        self.assertEqual(archived["value"], '=CONCAT("房门","打开")')
        self.assertEqual(archived["cachedValue"], "房门打开")
        self.assertTrue(any("Sheet!B3" in issue["text"] for issue in result["reviewIssues"]))
        self.assertTrue(any("Sheet!B4" in issue["text"] for issue in result["reviewIssues"]))

    def test_legacy_unknown_sheet_is_not_silently_ignored(self):
        workbook = openpyxl.Workbook()
        workbook.active.append(["镜号", "画面内容"])
        workbook.active.append([1, "已识别镜头"])
        workbook.create_sheet("续页").append(["另一个脚本块"])
        workbook.save(self.source)
        result = module.import_excel(self.source, self.folder / "project")
        self.assertEqual(len(result["shots"]), 1)
        self.assertTrue(any("续页" in issue["text"] for issue in result["reviewIssues"]))
        self.assertEqual(len(result["source"]["sheets"]), 2)

    def test_invalid_mapping_fails_before_destination_is_created(self):
        workbook = openpyxl.Workbook()
        workbook.active.append(["序号", "画面"])
        workbook.active.append([1, "抬眼"])
        workbook.save(self.source)
        valid = {"sheet": "Sheet", "headerRow": 1, "contentColumn": "B"}
        invalid = [{}, {"sheets": []}, {"sheets": [{**valid, "sheet": "不存在"}]},
                   {"sheets": [valid, valid]}, {"sheets": [{**valid, "headerRow": True}]},
                   {"sheets": [{**valid, "contentColumn": "ZZ"}]},
                   {"sheets": [{**valid, "numberColumn": "B"}]},
                   {"sheets": [{**valid, "extra": 1}]}]
        for index, mapping in enumerate(invalid):
            target = self.folder / f"invalid-{index}"
            with self.subTest(mapping=mapping), self.assertRaises(ValueError):
                module.import_excel(self.source, target, mapping)
            self.assertFalse(target.exists())

    def test_different_mapping_prevents_false_duplicate_sheet_drop(self):
        workbook = openpyxl.Workbook()
        workbook.active.append(["人物开门", "人物坐下"])
        copy = workbook.create_sheet("副本")
        copy.append(["人物开门", "人物坐下"])
        workbook.save(self.source)
        result = module.import_excel(self.source, self.folder / "project", {"sheets": [
            {"sheet": "Sheet", "headerRow": 0, "contentColumn": "A"},
            {"sheet": "副本", "headerRow": 0, "contentColumn": "B"},
        ]})
        self.assertEqual([shot["original"] for shot in result["shots"]], ["人物开门", "人物坐下"])
        self.assertEqual(result["source"]["duplicates"], [])

    def test_merged_content_with_distinct_numbers_reports_source_ambiguity(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["镜号", "画面内容"])
        sheet.append([1, "两人先后入门"])
        sheet.append([2, None])
        sheet.merge_cells("B2:B3")
        workbook.save(self.source)
        result = module.import_excel(self.source, self.folder / "project")
        self.assertEqual(len(result["shots"]), 1)
        self.assertEqual(result["source"]["sheets"][0]["cells"]["A3"]["value"], 2)
        self.assertTrue(any(issue["id"] == "merged-numbers-S001" for issue in result["reviewIssues"]))

    def test_cli_reads_explicit_mapping_file(self):
        workbook = openpyxl.Workbook()
        workbook.active.append(["主体", "镜头描述"])
        workbook.active.append(["妈妈", "妈妈站在门口"])
        workbook.save(self.source)
        path = self.folder / "mapping.json"
        path.write_text(json.dumps({"sheets": [{"sheet": "Sheet", "headerRow": 1, "contentColumn": "B"}]}, ensure_ascii=False), encoding="utf-8-sig")
        result = subprocess.run([sys.executable, str(ROOT / "scripts/import_excel.py"), str(self.source), str(self.folder / "cli-project"), "--mapping", str(path)], capture_output=True, text=True, check=True)
        summary = json.loads(result.stdout)
        self.assertEqual(summary["shots"], 1)
        self.assertEqual(summary["reviewIssues"], [])


if __name__ == "__main__":
    unittest.main()

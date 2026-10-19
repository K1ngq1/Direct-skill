import base64
import importlib.util
import io
import json
import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl.drawing.image import Image

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("import_excel", ROOT / "scripts/import_excel.py")
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aXioAAAAASUVORK5CYII=")


class ImportTest(unittest.TestCase):
    def test_alternate_headers_keep_dialogue_and_notes(self):
        (ROOT.parent / "artifacts").mkdir(exist_ok=True)
        folder = Path(tempfile.mkdtemp(prefix="excel-columns-", dir=ROOT.parent / "artifacts"))
        source = folder / "columns.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(["镜号", "画面内容", "台词/旁白", "音乐/音效", "备注"])
        sheet.append([1, "房门打开", "妈妈：起来吃饭", "风声", "手持锅铲"])
        workbook.save(source)
        result = module.import_excel(source, folder / "project")
        self.assertEqual(result["promptMode"], "static")
        self.assertEqual(result["staticPromptFormat"], "concise-v1")
        shot = result["shots"][0]
        self.assertEqual(shot["original"], "房门打开")
        self.assertEqual(shot["originalFields"][0], {"name": "台词/旁白", "value": "妈妈：起来吃饭", "cell": "C2"})
        self.assertIn("手持锅铲", (folder / "project/sources/source.md").read_text(encoding="utf-8"))

    def test_duplicate_cells_images_and_source_are_preserved(self):
        artifacts = ROOT.parent / "artifacts"
        artifacts.mkdir(exist_ok=True)
        folder = Path(tempfile.mkdtemp(prefix="excel-verification-", dir=artifacts))
        source = folder / "输入 样例.xlsx"
        workbook = openpyxl.Workbook()
        for sheet in [workbook.active, workbook.create_sheet("副本")]:
            sheet["A1"], sheet["B1"] = "故事概览", "中文脚本"
            sheet.merge_cells("B1:C1")
            sheet["A2"] = "人物参考"
            image = Image(io.BytesIO(PNG)); sheet.add_image(image, "B2")
            sheet["A4"], sheet["B4"] = "分镜号", "分镜内容"
            sheet["A5"], sheet["B5"] = 1, "人物抬眼"
        workbook.save(source)
        original = source.read_bytes()
        result = module.import_excel(source, folder / "project")
        self.assertEqual(len(result["shots"]), 1)
        self.assertEqual(len(result["references"]), 1)
        self.assertEqual(result["source"]["duplicates"][0]["sameAs"], "Sheet")
        self.assertEqual(result["source"]["sheets"][0]["mergedRanges"], ["B1:C1"])
        self.assertEqual(result["references"][0]["sourceRef"]["anchorCell"], "B2")
        self.assertEqual(result["shots"][0]["sourceRef"]["contentCell"], "B5")
        self.assertEqual(result["shots"][0]["variants"], [])
        self.assertEqual(source.read_bytes(), original)
        self.assertIn("../references/R001.png", (folder / "project/sources/source.md").read_text(encoding="utf-8"))
        with self.assertRaises(ValueError):
            module.import_excel(source, folder / "project")


if __name__ == "__main__":
    unittest.main()

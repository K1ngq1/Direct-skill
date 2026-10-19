"""Archive XLSX text, cell locations and anchored images; never invent shot content."""
import argparse
import hashlib
import json
import math
import re
import shutil
from pathlib import Path

import openpyxl


def digest(data):
    return hashlib.sha256(data).hexdigest()


def cell_value(value):
    return value if isinstance(value, (str, int, float, bool, type(None))) else str(value)


def normalize_mapping(mapping, workbook):
    if mapping is None:
        return None
    if not isinstance(mapping, dict) or set(mapping) != {"sheets"} or not isinstance(mapping["sheets"], list) or not mapping["sheets"]:
        raise ValueError("mapping 须为非空的 {sheets: [{sheet, headerRow, contentColumn, numberColumn?, durationColumn?}]}。")
    result = {}
    for item in mapping["sheets"]:
        allowed = {"sheet", "headerRow", "contentColumn", "numberColumn", "durationColumn"}
        if not isinstance(item, dict) or set(item) - allowed or not {"sheet", "headerRow", "contentColumn"} <= set(item):
            raise ValueError("每个工作表映射须提供 sheet/headerRow/contentColumn，只接受约定字段。")
        name = item["sheet"]
        if not isinstance(name, str) or name not in workbook.sheetnames or name in result:
            raise ValueError("映射工作表必须存在且不重复。")
        header = item["headerRow"]
        if type(header) is not int or not 0 <= header <= workbook[name].max_row:
            raise ValueError(f"{name} 的 headerRow 须为有效行号；无表头使用 0。")
        normalized = {"sheet": name, "headerRow": header}
        for key in ("contentColumn", "numberColumn", "durationColumn"):
            value = item.get(key)
            if value is None and key != "contentColumn":
                normalized[key] = None
                continue
            if isinstance(value, str) and re.fullmatch(r"[A-Za-z]+", value):
                value = openpyxl.utils.column_index_from_string(value.upper())
            if type(value) is not int or not 1 <= value <= workbook[name].max_column:
                raise ValueError(f"{name} 的 {key} 须为已存在的 Excel 列字母或从 1 开始的整数。")
            normalized[key] = value
        columns = [normalized[key] for key in ("contentColumn", "numberColumn", "durationColumn") if normalized[key] is not None]
        if len(columns) != len(set(columns)):
            raise ValueError(f"{name} 的画面、镜号、时长列不能重复。")
        result[name] = normalized
    return result


def source_cell(sheet, row, column):
    """Read merged cells from their anchor without losing the covered source rows."""
    cell = sheet.cell(row, column)
    for merged in sheet.merged_cells.ranges:
        if cell.coordinate in merged:
            return sheet.cell(merged.min_row, merged.min_col), merged
    return cell, None


def import_excel(source, destination, mapping=None):
    source, destination = Path(source).resolve(), Path(destination).resolve()
    if source.suffix.lower() != ".xlsx":
        raise ValueError("第一版导入器只支持 XLSX；其他格式请先由 Agent 转换并核对。")
    if destination.exists():
        raise ValueError("目标目录已存在，请指定新的项目目录，避免覆盖已有资料。")
    workbook = openpyxl.load_workbook(source, data_only=False)
    cached_workbook = openpyxl.load_workbook(source, data_only=True)
    mappings = normalize_mapping(mapping, workbook)
    review_issues = []

    def issue(issue_id, message):
        review_issues.append({"id": issue_id, "text": message, "resolved": False})

    def value_of(sheet_name, cell):
        value = cached_workbook[sheet_name][cell.coordinate] if cell.data_type == "f" else cell
        return None if value.data_type == "e" else value.value

    sheets = []
    for sheet in workbook:
        cells = {c.coordinate: {"value": cell_value(c.value), "type": c.data_type}
                 for row in sheet for c in row if c.value is not None}
        for coordinate, archived in cells.items():
            if archived["type"] == "f":
                cached = cached_workbook[sheet.title][coordinate]
                archived.update(cachedValue=cell_value(cached.value), cachedType=cached.data_type)
                if cached.value is None or cached.data_type == "e":
                    issue(f"formula-cache-{len(review_issues) + 1}", f"{sheet.title}!{coordinate} 的公式缺少有效缓存结果，未作为画面正文或参数；请由 Agent 核对原稿，不能把公式当成内容。")
            elif archived["type"] == "e":
                issue(f"cell-error-{len(review_issues) + 1}", f"{sheet.title}!{coordinate} 含 Excel 错误值，已归档但未作为画面正文或参数。")
        images = []
        for image in sheet._images:
            # openpyxl has no public anchored-image extraction API.
            anchor = image.anchor._from
            data = image._data()
            images.append({"row": anchor.row + 1, "column": anchor.col + 1,
                           "width": image.width, "height": image.height,
                           "sha256": digest(data), "format": image.format, "data": data})
        sheets.append({"name": sheet.title, "cells": cells,
                       "mergedRanges": [str(r) for r in sheet.merged_cells.ranges], "images": images})
    destination.mkdir(parents=True)
    (destination / "sources").mkdir()
    (destination / "references").mkdir()
    shutil.copyfile(source, destination / "sources/source.xlsx")
    seen, references, shots, duplicates, blocks = {}, [], [], [], []
    for sheet in sheets:
        comparable = {"cells": sheet["cells"], "mergedRanges": sheet["mergedRanges"],
                      "images": [{k: v for k, v in im.items() if k != "data"} for im in sheet["images"]]}
        if mappings is not None:
            comparable["mapping"] = {k: v for k, v in mappings.get(sheet["name"], {}).items() if k != "sheet"}
        signature = digest(json.dumps(comparable, sort_keys=True, ensure_ascii=False).encode("utf-8"))
        if signature in seen:
            duplicates.append({"sheet": sheet["name"], "sameAs": seen[signature]})
            continue
        seen[signature] = sheet["name"]
        blocks.append({"sheet": sheet["name"], "cells": sheet["cells"], "mergedRanges": sheet["mergedRanges"]})
        ws = workbook[sheet["name"]]
        header, duration_col = None, None
        if mappings is not None:
            selection = mappings.get(sheet["name"])
            if selection:
                header_row, number_col, content_col, duration_col = (selection[k] for k in ("headerRow", "numberColumn", "contentColumn", "durationColumn"))
                header = (header_row, number_col, content_col)
            elif sheet["cells"] or sheet["images"]:
                issue(f"unmapped-sheet-{len(review_issues) + 1}", f"工作表“{sheet['name']}”未在映射中指定，全部单元格、合并关系和图片已归档，尚未转换为分镜；请 Agent 核对，不能忽略该表。")
        else:
            for row in ws:
                columns = {str(value_of(ws.title, c)).strip(): c.column for c in row if value_of(ws.title, c) is not None}
                number_col = next((columns[k] for k in ("分镜号", "镜号") if k in columns), None)
                content_col = next((columns[k] for k in ("分镜内容", "画面内容") if k in columns), None)
                if number_col and content_col:
                    header = (row[0].row, number_col, content_col)
                    break
            if header is None and (sheet["cells"] or sheet["images"]):
                issue(f"unrecognized-sheet-{len(review_issues) + 1}", f"工作表“{sheet['name']}”未识别到分镜表头，内容已完整归档；请 Agent 提供显式列映射或核对用途。")
        row_to_shot = {}
        if header:
            header_row, number_col, content_col = header
            columns = [(str(value_of(ws.title, ws.cell(header_row, col)) or openpyxl.utils.get_column_letter(col)), col)
                       if header_row else (openpyxl.utils.get_column_letter(col), col)
                       for col in range(1, ws.max_column + 1)]
            used_contents = set()
            for row_number in range(header_row + 1, ws.max_row + 1):
                content, content_range = source_cell(ws, row_number, content_col)
                if content.row <= header_row or content.value is None or content.coordinate in used_contents:
                    continue
                used_contents.add(content.coordinate)
                number = source_cell(ws, row_number, number_col)[0] if number_col else None
                number_value = value_of(ws.title, number) if number else None
                content_value = value_of(ws.title, content)
                end_row = content_range.max_row if content_range else row_number
                shot_id = f"S{len(shots) + 1:03d}"
                if number_col and end_row > row_number:
                    covered_numbers = {str(value_of(ws.title, source_cell(ws, r, number_col)[0]))
                                       for r in range(row_number, end_row + 1)
                                       if value_of(ws.title, source_cell(ws, r, number_col)[0]) is not None}
                    if len(covered_numbers) > 1:
                        issue(f"merged-numbers-{shot_id}", f"{ws.title}!{content_range} 的合并画面对应多个原镜号，暂以一个内容块归档；各原镜号仍在原单元格中，请 Agent 判读拆分关系。")
                for covered_row in range(row_number, end_row + 1):
                    row_to_shot[covered_row] = shot_id
                source_ref = {"sheet": ws.title, "numberCell": number.coordinate if number else None,
                              "contentCell": content.coordinate}
                if content_range:
                    source_ref["contentRange"] = str(content_range)
                fields, used_fields = [], set()
                for covered_row in range(row_number, end_row + 1):
                    for label, col in columns:
                        field, _ = source_cell(ws, covered_row, col)
                        if col in (number_col, content_col) or field.coordinate == content.coordinate or field.coordinate in used_fields:
                            continue
                        used_fields.add(field.coordinate)
                        value = value_of(ws.title, field)
                        if value is not None:
                            fields.append({"name": label, "value": str(value), "cell": field.coordinate})
                duration = None
                if duration_col:
                    duration_cell = source_cell(ws, row_number, duration_col)[0]
                    source_ref["durationCell"] = duration_cell.coordinate
                    value = value_of(ws.title, duration_cell)
                    if value is not None:
                        match = re.fullmatch(r"\s*(\d+(?:\.\d+)?)\s*(?:秒|s)?\s*", str(value), re.IGNORECASE)
                        if match and math.isfinite(float(match[1])) and float(match[1]) > 0:
                            duration = float(match[1])
                        else:
                            issue(f"duration-{shot_id}", f"{ws.title}!{duration_cell.coordinate} 的时长不是明确的正数秒，已保留原值；请 Agent 判读单位，未自动换算。")
                shots.append({"id": shot_id, "sourceNumber": "" if number_value is None else str(number_value),
                              "original": "" if content_value is None else str(content_value),
                              "originalFields": fields, "sourceRef": source_ref,
                              "originalDuration": duration, "suggestedDuration": None, "durationOrigin": "source" if duration is not None else "unknown",
                              "characters": [], "scene": "待分析", "tags": [], "purpose": "", "variants": [],
                              "selectedVariantId": None, "approval": None, "status": "unplanned"})
        for im in sheet["images"]:
            reference_id = f"R{len(references) + 1:03d}"
            relative_path = f"references/{reference_id}.{im['format']}"
            (destination / relative_path).write_bytes(im["data"])
            row_label = value_of(sheet["name"], workbook[sheet["name"]].cell(im["row"], 1))
            references.append({"id": reference_id, "path": relative_path, "sha256": im["sha256"],
                               "sourceRef": {"sheet": sheet["name"], "anchorCell":
                                             f"{openpyxl.utils.get_column_letter(im['column'])}{im['row']}"},
                               "label": str(row_label or "嵌入图片"), "shotId": row_to_shot.get(im["row"]),
                               "width": im["width"], "height": im["height"], "role": "unclassified",
                               "visualStatus": "unreviewed", "observations": "", "usage": "", "uncertainties": ""})
    manifest = {"originalName": source.name, "path": "sources/source.xlsx", "sha256": digest(source.read_bytes()),
                "sheets": blocks, "duplicates": duplicates}
    if mappings is not None:
        manifest["mapping"] = {"sheets": list(mappings.values())}
    overview = next((str(value_of(s["sheet"], workbook[s["sheet"]]["B1"])) for s in blocks
                     if "B1" in s["cells"] and value_of(s["sheet"], workbook[s["sheet"]]["B1"]) is not None), "")
    project = {"schemaVersion": 1, "revision": 1, "promptMode": "static", "staticPromptFormat": "concise-v1", "title": source.stem, "overview": overview,
               "source": manifest, "aspectRatio": "待确认", "targetDuration": None,
               "tags": [], "tagOptions": [], "sharedConstraints": [], "reviewIssues": review_issues,
               "visualReview": {"status": "unreviewed", "reviewer": "", "evidence": ""},
               "shots": shots, "references": references, "comments": [],
               "lastChange": {"actor": "importer", "reason": "读取原件；尚未进行导演分析"}}
    if not shots:
        project["reviewIssues"].append({"id": "missing-shot-header", "text": "未提取到分镜；请 Agent 核对表头、列映射及正文范围。原单元格已完整归档。", "resolved": False})
    if any(c["type"] == "f" for s in blocks for c in s["cells"].values()):
        project["reviewIssues"].append({"id": "formulas", "text": "原稿含公式；归档保留公式及缓存值，正文只读取有效缓存，不执行或重新计算公式，请核对缓存是否为最新。", "resolved": False})
    write_json(destination / "project.json", project)
    write_json(destination / "sources/manifest.json", manifest)
    markdown = [f"# 原始脚本：{source.name}", "", "原始单元格与合并关系见 sources/manifest.json。", "", overview, ""]
    for shot in shots:
        label = f"原镜号 {shot['sourceNumber']}" if shot["sourceNumber"] else f"内部镜头 {shot['id']}（原稿未提供镜号）"
        markdown += [f"## {label}", f"来源：{shot['sourceRef']['sheet']} / {shot['sourceRef']['contentCell']}", "", shot["original"], ""]
        for field in shot["originalFields"]:
            markdown += [f"{field['name']}（{field['cell']}）：{field['value']}", ""]
    for ref in references:
        markdown += [f"## {ref['id']}：{ref['label']}", f"来源：{ref['sourceRef']['sheet']} / {ref['sourceRef']['anchorCell']}",
                     f"![{ref['label']}](../{ref['path']})", ""]
    (destination / "sources/source.md").write_text("\n".join(markdown), encoding="utf-8")
    return project


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source")
    parser.add_argument("destination")
    parser.add_argument("--mapping", help="UTF-8 JSON 列映射文件；不提供时按已有双表头规则识别。")
    args = parser.parse_args()
    mapping = json.loads(Path(args.mapping).read_text(encoding="utf-8-sig")) if args.mapping else None
    result = import_excel(args.source, args.destination, mapping)
    print(json.dumps({"shots": len(result["shots"]), "references": len(result["references"]),
                      "duplicateSheets": result["source"]["duplicates"], "reviewIssues": result["reviewIssues"]}, ensure_ascii=True))
